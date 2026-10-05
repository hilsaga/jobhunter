"""Open, readable layout for the master CV PDF and Word file.

The text in docs/master_cv.txt stays the source. This module only changes
how that text is set on the page.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import HRFlowable, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from src.textutil import escape_xml

_INK = HexColor("#1C1917")
_ACCENT = HexColor("#2F4F46")
_MUTED = HexColor("#6B6560")
_RULE = HexColor("#E4DDD4")
_CHIP = HexColor("#F4F0E8")
_PAPER = HexColor("#FBF9F6")
_ROW = HexColor("#F7F4EF")

_STACK_LABELS = (
    "AI and architecture",
    "Leadership and delivery",
    "Analysis and deployment",
    "Languages",
    "Platforms and data",
    "Markets",
)
_SECTIONS = {
    "PROFESSIONAL SUMMARY": "summary",
    "KEYWORDS": "keywords",
    "CORE COMPETENCIES & TECHNICAL STACK": "stack",
    "PROFESSIONAL EXPERIENCE": "experience",
    "EDUCATION": "education",
}
_JOB = re.compile(
    r"^(?P<org>.+?)\s*\|\s*(?P<place>.+?)\s+-\s+(?P<title>.+?)\s*\((?P<dates>[^)]+)\)\s*$"
)
_DEGREE = re.compile(r"^(?P<degree>.+?)\s*\|\s*(?P<school>.+?)\s*\((?P<dates>[^)]+)\)\s*$")


@dataclass
class JobBlock:
    organization: str
    place: str
    title: str
    dates: str
    bullets: list[str] = field(default_factory=list)


@dataclass
class MasterCv:
    name: str
    contact: str
    languages: str
    roles: list[str]
    summary: str
    keywords: list[str]
    stack: list[str]
    jobs: list[JobBlock]
    education: list[str]


def write_master_documents(text: str, pdf_path: Path, docx_path: Path) -> None:
    """Write the styled PDF and Word copies of the master CV."""
    document = parse_master_cv(text)
    _write_pdf(document, pdf_path)
    _write_docx(document, docx_path)


def parse_master_cv(text: str) -> MasterCv:
    """Split the master CV text into the blocks the layout uses."""
    lines = [line.rstrip() for line in text.splitlines()]
    header: list[str] = []
    buckets: dict[str, list[str]] = {name: [] for name in _SECTIONS.values()}
    current = ""
    for line in lines:
        section = _SECTIONS.get(line.strip())
        if section:
            current = section
            continue
        if not current:
            if line.strip():
                header.append(line.strip())
            continue
        buckets[current].append(line)

    roles = [part.strip() for part in header[3].split("|")] if len(header) > 3 else []
    jobs: list[JobBlock] = []
    for line in buckets["experience"]:
        stripped = line.strip()
        if not stripped:
            continue
        if stripped.startswith("- "):
            if jobs:
                jobs[-1].bullets.append(stripped[2:].strip())
            continue
        match = _JOB.match(stripped)
        if match:
            jobs.append(
                JobBlock(
                    organization=match.group("org").strip(),
                    place=match.group("place").strip(),
                    title=match.group("title").strip(),
                    dates=match.group("dates").strip(),
                    bullets=[],
                )
            )
    keywords: list[str] = []
    for line in buckets["keywords"]:
        keywords.extend(part.strip() for part in line.split(",") if part.strip())
    return MasterCv(
        name=header[0] if header else "",
        contact=header[1] if len(header) > 1 else "",
        languages=header[2] if len(header) > 2 else "",
        roles=roles,
        summary=" ".join(line.strip() for line in buckets["summary"] if line.strip()),
        keywords=keywords,
        stack=[line.strip() for line in buckets["stack"] if line.strip()],
        jobs=jobs,
        education=[line.strip() for line in buckets["education"] if line.strip()],
    )


def _fonts() -> dict[str, str]:
    pairs = {
        "serif": ("/System/Library/Fonts/Supplemental/Georgia.ttf", "Times-Roman"),
        "serif-bold": ("/System/Library/Fonts/Supplemental/Georgia Bold.ttf", "Times-Bold"),
        "serif-italic": ("/System/Library/Fonts/Supplemental/Georgia Italic.ttf", "Times-Italic"),
    }
    chosen: dict[str, str] = {}
    for name, (path, fallback) in pairs.items():
        if Path(path).is_file():
            pdfmetrics.registerFont(TTFont(f"MasterCv-{name}", path))
            chosen[name] = f"MasterCv-{name}"
        else:
            chosen[name] = fallback
    chosen["sans"] = "Helvetica"
    chosen["sans-bold"] = "Helvetica-Bold"
    return chosen


def _write_pdf(document: MasterCv, path: Path) -> None:
    fonts = _fonts()
    styles = _pdf_styles(fonts)
    story: list[object] = [
        Paragraph(escape_xml(document.name), styles["name"]),
        Spacer(1, 8),
    ]
    if document.contact.strip():
        story.append(Paragraph(_dotted(document.contact), styles["contact"]))
        story.append(Spacer(1, 3))
    if document.languages.strip():
        story.append(Paragraph(escape_xml(document.languages), styles["meta"]))
    story.append(Spacer(1, 12))
    for role in document.roles:
        story.append(Paragraph(escape_xml(role), styles["role"]))
    story.append(Spacer(1, 8))
    story.append(HRFlowable(width="100%", thickness=0.8, color=_ACCENT, spaceBefore=2, spaceAfter=6))
    if document.summary.strip():
        story.extend(_pdf_section("Professional summary", [Paragraph(escape_xml(document.summary), styles["body"])], styles))
    if document.keywords:
        story.extend(_pdf_section("Keywords", [_keyword_table(document.keywords, styles)], styles))
    if document.stack:
        story.append(
            KeepTogether(
                _pdf_section("Core competencies and technical stack", [_stack_table(document, styles)], styles)
            )
        )
    experience: list[object] = []
    for job in document.jobs:
        experience.append(KeepTogether(_job_flowables(job, styles)))
    story.extend(_pdf_section("Professional experience", experience, styles))
    education: list[object] = []
    for line in document.education:
        block = _education_block(line, styles)
        if isinstance(block, list):
            education.extend(block)
        else:
            education.append(block)
    story.extend(_pdf_section("Education", education, styles))

    path.parent.mkdir(parents=True, exist_ok=True)
    pdf = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=58,
        rightMargin=58,
        topMargin=46,
        bottomMargin=42,
        title=document.name or "Master CV",
        author=document.name or "Master CV",
    )
    pdf.build(story, onFirstPage=_paint, onLaterPages=_paint)


def _pdf_styles(fonts: dict[str, str]) -> dict[str, ParagraphStyle]:
    return {
        "name": ParagraphStyle(
            "MasterName",
            fontName=fonts["serif-bold"],
            fontSize=23,
            leading=28,
            textColor=_INK,
            alignment=TA_LEFT,
            spaceAfter=0,
        ),
        "contact": ParagraphStyle(
            "MasterContact",
            fontName=fonts["sans"],
            fontSize=10,
            leading=14,
            textColor=_MUTED,
        ),
        "meta": ParagraphStyle(
            "MasterMeta",
            fontName=fonts["sans"],
            fontSize=10,
            leading=14,
            textColor=_MUTED,
        ),
        "role": ParagraphStyle(
            "MasterRole",
            fontName=fonts["serif-italic"],
            fontSize=12,
            leading=17,
            textColor=_ACCENT,
            spaceBefore=1,
            spaceAfter=1,
        ),
        "section": ParagraphStyle(
            "MasterSection",
            fontName=fonts["serif-bold"],
            fontSize=12,
            leading=16,
            textColor=_ACCENT,
            spaceBefore=0,
            spaceAfter=2,
        ),
        "body": ParagraphStyle(
            "MasterBody",
            fontName=fonts["sans"],
            fontSize=10.5,
            leading=17,
            textColor=_INK,
            spaceAfter=2,
        ),
        "label": ParagraphStyle(
            "MasterLabel",
            fontName=fonts["sans-bold"],
            fontSize=8.5,
            leading=12,
            textColor=_ACCENT,
        ),
        "chip": ParagraphStyle(
            "MasterChip",
            fontName=fonts["sans"],
            fontSize=8.5,
            leading=12,
            textColor=_INK,
            alignment=TA_LEFT,
        ),
        "org": ParagraphStyle(
            "MasterOrg",
            fontName=fonts["serif-bold"],
            fontSize=12,
            leading=16,
            textColor=_INK,
        ),
        "dates": ParagraphStyle(
            "MasterDates",
            fontName=fonts["sans"],
            fontSize=9.5,
            leading=13,
            textColor=_MUTED,
            alignment=TA_RIGHT,
        ),
        "title": ParagraphStyle(
            "MasterTitle",
            fontName=fonts["sans"],
            fontSize=10.5,
            leading=15,
            textColor=_ACCENT,
            spaceBefore=1,
            spaceAfter=4,
        ),
        "bullet": ParagraphStyle(
            "MasterBullet",
            fontName=fonts["sans"],
            fontSize=10.5,
            leading=16,
            textColor=_INK,
            leftIndent=14,
            bulletIndent=0,
            bulletFontName=fonts["sans"],
            bulletFontSize=10.5,
            spaceBefore=1,
            spaceAfter=3,
        ),
        "degree": ParagraphStyle(
            "MasterDegree",
            fontName=fonts["sans-bold"],
            fontSize=10.5,
            leading=15,
            textColor=_INK,
            spaceBefore=4,
        ),
        "school": ParagraphStyle(
            "MasterSchool",
            fontName=fonts["sans"],
            fontSize=10,
            leading=14,
            textColor=_MUTED,
            spaceAfter=2,
        ),
    }


def _pdf_section(title: str, blocks: list[object], styles: dict[str, ParagraphStyle]) -> list[object]:
    return [
        Spacer(1, 16),
        Paragraph(escape_xml(title), styles["section"]),
        HRFlowable(width="100%", thickness=0.4, color=_RULE, spaceBefore=1, spaceAfter=8),
        *blocks,
    ]


def _keyword_table(keywords: list[str], styles: dict[str, ParagraphStyle]) -> Table:
    columns = 3
    cells = [Paragraph(escape_xml(word), styles["chip"]) for word in keywords]
    while len(cells) % columns:
        cells.append(Paragraph("", styles["chip"]))
    rows = [cells[index : index + columns] for index in range(0, len(cells), columns)]
    table = Table(rows, colWidths=[158, 158, 158], hAlign="LEFT")
    style_commands: list[tuple] = [
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 5),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ("BACKGROUND", (0, 0), (-1, -1), _CHIP),
        ("ROWBACKGROUNDS", (0, 0), (-1, -1), [_CHIP, _CHIP]),
        ("LINEBEFORE", (1, 0), (-1, -1), 5, _PAPER),
        ("LINEABOVE", (0, 1), (-1, -1), 5, _PAPER),
    ]
    table.setStyle(TableStyle(style_commands))
    return table


def _stack_table(document: MasterCv, styles: dict[str, ParagraphStyle]) -> Table:
    rows = []
    for index, line in enumerate(document.stack):
        label = _STACK_LABELS[index] if index < len(_STACK_LABELS) else "Also"
        items = "   ·   ".join(part.strip() for part in line.split(",") if part.strip())
        rows.append(
            [
                Paragraph(escape_xml(label), styles["label"]),
                Paragraph(escape_xml(items), styles["body"]),
            ]
        )
    table = Table(rows, colWidths=[128, 346], hAlign="LEFT")
    table.splitByRow = 0
    commands: list[tuple] = [
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (0, -1), 8),
        ("RIGHTPADDING", (0, 0), (0, -1), 8),
        ("LEFTPADDING", (1, 0), (1, -1), 10),
        ("RIGHTPADDING", (1, 0), (1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
        ("BACKGROUND", (0, 0), (0, -1), _ROW),
        ("LINEBELOW", (0, 0), (-1, -2), 0.3, _RULE),
    ]
    table.setStyle(TableStyle(commands))
    return table


def _job_flowables(job: JobBlock, styles: dict[str, ParagraphStyle]) -> list[object]:
    org_line = job.organization
    if job.place:
        org_line = f"{job.organization}   ·   {job.place}"
    heading = Table(
        [[
            Paragraph(escape_xml(org_line), styles["org"]),
            Paragraph(escape_xml(job.dates), styles["dates"]),
        ]],
        colWidths=[360, 114],
        hAlign="LEFT",
    )
    heading.setStyle(
        TableStyle(
            [
                ("VALIGN", (0, 0), (-1, -1), "BOTTOM"),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                ("TOPPADDING", (0, 0), (-1, -1), 0),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
            ]
        )
    )
    blocks: list[object] = [
        Spacer(1, 12),
        heading,
        Paragraph(escape_xml(job.title), styles["title"]),
    ]
    blocks.extend(Paragraph(escape_xml(bullet), styles["bullet"], bulletText="·") for bullet in job.bullets)
    return blocks


def _education_block(line: str, styles: dict[str, ParagraphStyle]) -> object:
    match = _DEGREE.match(line)
    if not match:
        return Paragraph(escape_xml(line), styles["body"])
    return [
        Paragraph(escape_xml(match.group("degree")), styles["degree"]),
        Paragraph(escape_xml(f"{match.group('school')}   ·   {match.group('dates')}"), styles["school"]),
    ]


def _paint(canvas, _doc) -> None:  # type: ignore[no-untyped-def]
    canvas.saveState()
    canvas.setFillColor(_PAPER)
    canvas.rect(0, 0, A4[0], A4[1], fill=1, stroke=0)
    canvas.setFillColor(_ACCENT)
    canvas.rect(0, A4[1] - 8, A4[0], 8, fill=1, stroke=0)
    canvas.restoreState()


def _dotted(value: str) -> str:
    parts = [escape_xml(part.strip()) for part in value.split("|") if part.strip()]
    return "    ·    ".join(parts)


def _write_docx(document: MasterCv, path: Path) -> None:
    word = Document()
    section = word.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(16)
    section.bottom_margin = Mm(16)
    _set_run_font(word.styles["Normal"], "Calibri", 11)

    name = word.add_paragraph()
    name.paragraph_format.space_after = Pt(4)
    _run(name, document.name, font="Georgia", size=22, bold=True, color="1C1917")

    contact = word.add_paragraph()
    contact.paragraph_format.space_before = Pt(2)
    contact.paragraph_format.space_after = Pt(0)
    _run(contact, "   ·   ".join(part.strip() for part in document.contact.split("|")), size=10.5, color="6B6560")

    languages = word.add_paragraph()
    languages.paragraph_format.space_before = Pt(1)
    languages.paragraph_format.space_after = Pt(8)
    _run(languages, document.languages, size=10.5, color="6B6560")

    for role in document.roles:
        line = word.add_paragraph()
        line.paragraph_format.space_before = Pt(0)
        line.paragraph_format.space_after = Pt(1)
        line.paragraph_format.line_spacing = 1.15
        _run(line, role, font="Georgia", size=12, italic=True, color="2F4F46")

    _docx_section(word, "Professional summary")
    summary = word.add_paragraph()
    summary.paragraph_format.space_after = Pt(4)
    summary.paragraph_format.line_spacing = 1.2
    _run(summary, document.summary, size=11, color="1C1917")

    _docx_section(word, "Keywords")
    _docx_keywords(word, document.keywords)

    _docx_section(word, "Core competencies and technical stack")
    _docx_stack(word, document)

    _docx_section(word, "Professional experience")
    for job in document.jobs:
        _docx_job(word, job)

    _docx_section(word, "Education")
    for line in document.education:
        _docx_education(word, line)

    path.parent.mkdir(parents=True, exist_ok=True)
    word.save(path)


def _docx_section(word: Document, title: str) -> None:
    paragraph = word.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(16)
    paragraph.paragraph_format.space_after = Pt(6)
    _run(paragraph, title, font="Georgia", size=13, bold=True, color="2F4F46")
    _bottom_border(paragraph)


def _docx_keywords(word: Document, keywords: list[str]) -> None:
    columns = 3
    padded = list(keywords)
    while len(padded) % columns:
        padded.append("")
    table = word.add_table(rows=len(padded) // columns, cols=columns)
    _clear_table_borders(table)
    table.autofit = True
    for index, word_text in enumerate(padded):
        cell = table.cell(index // columns, index % columns)
        cell.text = ""
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(1)
        if word_text:
            _run(paragraph, word_text, size=9, color="1C1917")
            _shade_cell(cell, "F4F0E8")
        _cell_margins(cell)


def _docx_stack(word: Document, document: MasterCv) -> None:
    table = word.add_table(rows=len(document.stack), cols=2)
    _clear_table_borders(table)
    for index, line in enumerate(document.stack):
        label = _STACK_LABELS[index] if index < len(_STACK_LABELS) else "Also"
        items = "   ·   ".join(part.strip() for part in line.split(",") if part.strip())
        label_cell = table.cell(index, 0)
        value_cell = table.cell(index, 1)
        label_cell.text = ""
        value_cell.text = ""
        label_paragraph = label_cell.paragraphs[0]
        value_paragraph = value_cell.paragraphs[0]
        label_paragraph.paragraph_format.space_after = Pt(0)
        value_paragraph.paragraph_format.space_after = Pt(0)
        value_paragraph.paragraph_format.line_spacing = 1.15
        _run(label_paragraph, label, size=9, bold=True, color="2F4F46")
        _run(value_paragraph, items, size=10.5, color="1C1917")
        _shade_cell(label_cell, "F7F4EF")
        _cell_margins(label_cell)
        _cell_margins(value_cell)
    _set_column_widths(table, (Mm(42), Mm(128)))


def _docx_job(word: Document, job: JobBlock) -> None:
    table = word.add_table(rows=1, cols=2)
    _clear_table_borders(table)
    left = table.cell(0, 0)
    right = table.cell(0, 1)
    left.text = ""
    right.text = ""
    left.paragraphs[0].paragraph_format.space_before = Pt(10)
    right.paragraphs[0].paragraph_format.space_before = Pt(10)
    right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _run(left.paragraphs[0], f"{job.organization}   ·   {job.place}", font="Georgia", size=12, bold=True, color="1C1917")
    _run(right.paragraphs[0], job.dates, size=10, color="6B6560")
    _set_column_widths(table, (Mm(120), Mm(50)))

    title = word.add_paragraph()
    title.paragraph_format.space_before = Pt(1)
    title.paragraph_format.space_after = Pt(4)
    _run(title, job.title, size=11, color="2F4F46")
    for bullet in job.bullets:
        paragraph = word.add_paragraph()
        paragraph.paragraph_format.left_indent = Mm(5)
        paragraph.paragraph_format.space_before = Pt(1)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.line_spacing = 1.15
        _run(paragraph, f"·   {bullet}", size=11, color="1C1917")


def _docx_education(word: Document, line: str) -> None:
    match = _DEGREE.match(line)
    if not match:
        paragraph = word.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(4)
        paragraph.paragraph_format.line_spacing = 1.15
        _run(paragraph, line, size=11, color="1C1917")
        return
    degree = word.add_paragraph()
    degree.paragraph_format.space_before = Pt(8)
    degree.paragraph_format.space_after = Pt(0)
    _run(degree, match.group("degree"), size=11, bold=True, color="1C1917")
    school = word.add_paragraph()
    school.paragraph_format.space_before = Pt(0)
    school.paragraph_format.space_after = Pt(2)
    _run(school, f"{match.group('school')}   ·   {match.group('dates')}", size=10.5, color="6B6560")


def _run(
    paragraph,
    text: str,
    *,
    font: str = "Calibri",
    size: float = 11,
    bold: bool = False,
    italic: bool = False,
    color: str = "1C1917",
) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    r_fonts = run._element.rPr.rFonts
    r_fonts.set(qn("w:ascii"), font)
    r_fonts.set(qn("w:hAnsi"), font)
    r_fonts.set(qn("w:eastAsia"), font)


def _set_run_font(style, font: str, size: int) -> None:
    style.font.name = font
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    r_fonts = rpr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        rpr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), font)
    r_fonts.set(qn("w:hAnsi"), font)


def _bottom_border(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "E4DDD4")
    borders.append(bottom)
    p_pr.append(borders)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shade = OxmlElement("w:shd")
    shade.set(qn("w:val"), "clear")
    shade.set(qn("w:color"), "auto")
    shade.set(qn("w:fill"), fill)
    tc_pr.append(shade)


def _cell_margins(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement("w:tcMar")
    for edge, value in (("top", "70"), ("bottom", "70"), ("left", "90"), ("right", "90")):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")
        margins.append(node)
    tc_pr.append(margins)


def _clear_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "nil")
        node.set(qn("w:sz"), "0")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "auto")
        borders.append(node)
    tbl_pr.append(borders)


def _set_column_widths(table, widths: tuple) -> None:
    table.autofit = False
    table.allow_autofit = False
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(1, grid)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(int(width.twips)))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths, strict=True):
            cell.width = width
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width.twips)))
            tc_w.set(qn("w:type"), "dxa")

