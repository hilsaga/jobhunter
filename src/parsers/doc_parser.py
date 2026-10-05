"""Scan and parse resume and supporting documents in docs/."""

from __future__ import annotations

import csv
import io
import logging
from pathlib import Path
from urllib.parse import urlparse

from src.models import ParsedDocument, Site

logger = logging.getLogger("jobhunter.parser")

SUPPORTED_SUFFIXES = {".pdf", ".docx", ".txt"}


class DocParser:
    """Scan docs/ and extract text from every supported resume or base file."""

    def __init__(self, docs_dir: Path) -> None:
        self.docs_dir = docs_dir

    def scan(self) -> list[ParsedDocument]:
        if not self.docs_dir.is_dir():
            logger.warning("Docs directory does not exist: %s", self.docs_dir)
            return []

        files = sorted(
            (
                path
                for path in self.docs_dir.rglob("*")
                if path.is_file()
                and not path.name.startswith(".")
                and path.suffix.lower() in SUPPORTED_SUFFIXES
            ),
            key=lambda path: document_sort_key(str(path)),
        )
        documents: list[ParsedDocument] = []
        for path in files:
            try:
                text = self._read(path)
            except Exception:
                logger.exception("Failed to parse %s", path)
                continue
            cleaned = text.strip()
            if not cleaned:
                logger.warning("No text extracted from %s", path)
                continue
            if len(cleaned) > 150_000:
                cleaned = cleaned[:150_000]
                logger.info("Truncated %s to 150000 characters", path.name)
            documents.append(ParsedDocument(path=str(path), text=cleaned))
            logger.info("Parsed %s (%s characters)", path.name, len(cleaned))
        return documents

    def _read(self, path: Path) -> str:
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            return _read_pdf(path)
        if suffix == ".docx":
            return _read_docx(path)
        if not path.is_file():
            raise FileNotFoundError(path)
        return path.read_text(encoding="utf-8", errors="replace")


def _read_pdf(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(path)
    text = _read_pdfplumber(path)
    if text.strip():
        return text
    logger.info("pdfplumber returned no text for %s; trying pypdf", path.name)
    return _read_pypdf(path)


def _read_pdfplumber(path: Path) -> str:
    import pdfplumber

    chunks: list[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            chunks.append(page.extract_text() or "")
    return "\n".join(chunks)


def _read_pypdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    if not path.is_file():
        raise FileNotFoundError(path)
    from docx import Document

    document = Document(str(path))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def document_sort_key(path: str) -> tuple[int, int, str]:
    """Prefer the master CV text, then other CVs, then supporting documents."""
    file = Path(path)
    stem = file.stem.lower()
    if "cover" in stem:
        kind = 3
    elif "master_cv" in stem or stem in {"resume", "curriculum_vitae"} or stem.startswith("resume"):
        kind = 0
    elif stem == "cv" or stem.startswith("cv_") or stem.startswith("cv-"):
        kind = 1
    else:
        kind = 2
    if kind == 0:
        suffix_rank = {".txt": 0, ".docx": 1, ".pdf": 2}.get(file.suffix.lower(), 9)
    else:
        suffix_rank = {".pdf": 0, ".docx": 1, ".txt": 2}.get(file.suffix.lower(), 9)
    return (kind, suffix_rank, file.name.lower())


def load_sites(path: Path) -> list[Site]:
    """Load sites.csv, including a headerless file that contains only URLs."""
    if not path.is_file():
        raise FileNotFoundError(path)
    raw = path.read_text(encoding="utf-8-sig")
    lines = [
        line.strip()
        for line in raw.splitlines()
        if line.strip() and not line.strip().startswith("#")
    ]
    if not lines:
        return []
    if lines[0].lower().startswith("http://") or lines[0].lower().startswith("https://"):
        return _sites_from_urls(lines)

    reader = csv.DictReader(io.StringIO(raw))
    sites: list[Site] = []
    for row in reader:
        if row is None:
            continue
        lowered = {(key or "").strip().lower(): (value or "").strip() for key, value in row.items()}
        url = lowered.get("url") or lowered.get("link") or ""
        if not url.startswith("http://") and not url.startswith("https://"):
            logger.warning("Skipping site row without an http(s) URL: %s", row)
            continue
        name = lowered.get("site_name") or lowered.get("name") or urlparse(url).netloc or "site"
        location_param = lowered.get("location_filter_param") or lowered.get("location") or ""
        sites.append(Site(site_name=name, url=url, location_filter_param=location_param))
    return sites


def _sites_from_urls(lines: list[str]) -> list[Site]:
    sites: list[Site] = []
    for line in lines:
        if not line.lower().startswith("http://") and not line.lower().startswith("https://"):
            logger.warning("Skipping non-URL row: %s", line)
            continue
        host = urlparse(line).netloc or "site"
        sites.append(Site(site_name=host, url=line, location_filter_param=""))
    return sites
